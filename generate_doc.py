import sys
try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
except ImportError:
    print("python-docx is not installed yet")
    sys.exit(1)

doc = Document()

# Title
title = doc.add_heading('Style Guide: Aplikasi HikeIt', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

doc.add_paragraph('Dokumen ini merupakan panduan antarmuka (Style Guide) untuk aplikasi HikeIt. Panduan ini mencakup semua desain dan komponen yang digunakan dalam aplikasi untuk menjaga konsistensi UI/UX.')

# 1. Tautan Website
doc.add_heading('1. Informasi Proyek & Tautan', level=1)
p = doc.add_paragraph('Apakah tautan website perlu dicantumkan? ')
p.add_run('Sangat disarankan').bold = True
p.add_run(' untuk dicantumkan jika aplikasi sudah di-deploy (misalnya di Vercel), atau setidaknya cantumkan tautan repositori (GitHub) sebagai referensi bagi reviewer/dosen.')
doc.add_paragraph('Tautan Aplikasi: [Isi dengan link Vercel/Website Anda]')
doc.add_paragraph('Tautan Repositori: [Isi dengan link GitHub Anda]')

# 2. Pondasi Desain
doc.add_heading('2. Pondasi Desain (Design Foundation)', level=1)

doc.add_heading('A. Tema & Palet Warna', level=2)
doc.add_paragraph('Aplikasi HikeIt menggunakan gaya desain Dark Theme (Mode Gelap) sebagai tema utama, dengan aksen hijau (Emerald) sebagai identitas merek.')

colors = [
    ('Background Utama', '#0A0A0A', 'Digunakan untuk latar belakang seluruh halaman.'),
    ('Background Card/Surface', '#161616', 'Digunakan untuk latar belakang formulir, kartu, atau elemen yang menonjol dari latar.'),
    ('Border (Garis Tepi)', '#2E2E2E', 'Digunakan untuk pemisah antar elemen, pinggiran kartu, dan input.'),
    ('Primary (Aksen & Tombol)', '#059669 (Emerald-600)', 'Warna utama untuk tombol aksi, link, dan elemen interaktif. Hover state: #10B981 (Emerald-500).'),
    ('Teks Utama (Headings)', '#FFFFFF (Putih)', 'Digunakan untuk judul dan teks yang butuh penekanan tinggi.'),
    ('Teks Sekunder (Label/Subteks)', 'Gray-300 / Gray-400', 'Digunakan untuk label form dan deskripsi agar tidak terlalu kontras.'),
    ('Warna Status: Error', 'Merah', 'Background: red-900/30, Teks: red-400, Border: red-800. Digunakan untuk peringatan atau pesan error saat login.'),
    ('Aksen Sekunder (Warning/Highlight)', 'Amber/Kuning', 'Digunakan untuk label Paket Bundling (bg-amber-500/20, text-amber-400) dan ikon rating bintang (text-yellow-500).')
]

table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Elemen'
hdr_cells[1].text = 'Kode Warna (Hex/Tailwind)'
hdr_cells[2].text = 'Penggunaan'

for item, color, usage in colors:
    row_cells = table.add_row().cells
    row_cells[0].text = item
    row_cells[1].text = color
    row_cells[2].text = usage

doc.add_heading('B. Tipografi & Sudut (Border Radius)', level=2)
doc.add_paragraph('• Font Utama: Geist Sans (Sans-serif)')
doc.add_paragraph('• Card Border Radius: Sangat membulat (rounded-2xl)')
doc.add_paragraph('• Input & Tombol Border Radius: Membulat (rounded-xl)')
doc.add_paragraph('• Shadow: Bayangan tebal (shadow-2xl) untuk memisahkan kartu dari latar belakang.')

# 3. Komponen
doc.add_heading('3. Komponen Utama & Interaktif', level=1)

doc.add_heading('A. Tombol (Buttons)', level=2)
doc.add_paragraph('Tombol Utama (Primary Button):')
p = doc.add_paragraph('• Default State: Latar belakang Emerald-600, teks putih, sudut rounded-xl, padding merata.')
p = doc.add_paragraph('• Hover State: Latar belakang berubah menjadi Emerald-500.')
p = doc.add_paragraph('• Disabled/Loading State: Opasitas berkurang (opacity-50) dan teks bisa berubah menjadi "Memeriksa...".')

doc.add_heading('B. Form & Input', level=2)
doc.add_paragraph('• Default State: Latar belakang hitam (#0A0A0A), border #2E2E2E, teks putih.')
doc.add_paragraph('• Focus State (Saat diklik): Border menyala dengan Emerald-500 dan bayangan pinggir (ring-2 ring-emerald-500/50).')
doc.add_paragraph('• Password Toggle: Terdapat ikon mata di sisi kanan input password untuk menampilkan/menyembunyikan teks.')

doc.add_heading('C. Kartu (Card)', level=2)
doc.add_paragraph('Digunakan untuk membungkus form login/register atau profil. Memiliki background #161616, border tipis, dan bayangan 2XL agar terlihat melayang di atas background utama.')

doc.add_heading('D. Notifikasi Error (Alerts)', level=2)
doc.add_paragraph('Kotak peringatan jika login gagal. Berupa kotak memanjang (rounded-lg) dengan background merah transparan dan teks merah terang, ditempatkan di atas form input.')

doc.add_heading('E. Antarmuka Peta (MapSelector)', level=2)
doc.add_paragraph('Menampilkan komponen peta interaktif. Memiliki kontrol Zoom In/Out dan kustom penanda lokasi (Pin Marker) yang disesuaikan dengan warna merek (Emerald).')

doc.add_heading('F. Navigasi & Dropdown Menu', level=2)
doc.add_paragraph('• Navbar: Menggunakan properti sticky (sticky top-0) agar selalu tampil di atas, latar belakang #161616, dan garis bawah border-[#2E2E2E].')
doc.add_paragraph('• Dropdown Profile: Muncul sebagai popup dengan posisi absolut (absolute right-0), border-[#2E2E2E], dan bayangan tebal (shadow-2xl z-50).')

doc.add_heading('G. Tabel Data (Data Table)', level=2)
doc.add_paragraph('Digunakan di dashboard mitra untuk menampilkan daftar alat. Memiliki desain header gelap (bg-[#0A0A0A]) dengan teks uppercase yang kecil (text-xs). Tiap baris tabel dipisahkan oleh garis (border-b border-[#2E2E2E]) dan memiliki efek hover terang (hover:bg-[#2E2E2E]/50).')

doc.add_heading('H. Modal / Dialog (Pop-up)', level=2)
doc.add_paragraph('Digunakan untuk konfirmasi atau form tambahan seperti Edit Alat, Balas Ulasan, dan Buat Bundling. Modal diletakkan di tengah layar dengan latar belakang gelap/redup (overlay) dan box konten di atasnya.')

doc.add_heading('I. Badges & Label Status', level=2)
doc.add_paragraph('Label kecil (teks ukuran text-[10px] atau text-xs) dengan latar transparan yang diblur (backdrop-blur-sm bg-[#0A0A0A]/80). Biasa digunakan untuk menampilkan kategori alat, teks diskon, atau penanda status paket bundling di pojok kartu.')

doc.add_heading('J. Status Loading & Animasi', level=2)
doc.add_paragraph('• Spinner Loading: Lingkaran berputar (animate-spin) dengan border Emerald (border-4 border-emerald-500 border-t-transparent) berukuran w-8 h-8.')
doc.add_paragraph('• Animasi Hover (Micro-animation): Gambar alat akan sedikit membesar ketika kursor diarahkan ke kartu (group-hover:scale-105 transition-transform duration-300).')

doc.add_heading('K. Komponen Rating & Ulasan (Review)', level=2)
doc.add_paragraph('Menampilkan komentar pelanggan dengan teks yang dimiringkan (italic) serta jajaran ikon bintang menggunakan pewarnaan text-yellow-500 untuk mencerminkan nilai rating.')

doc.add_page_break()
doc.add_heading('Lampiran: Screenshot Komponen', level=1)
doc.add_paragraph('[CATATAN UNTUK ANDA: Silakan masukkan screenshot halaman Login (page.tsx) dan MapSelector Anda di sini sebagai bukti visual implementasi desain]')

doc.save('Style_Guide_HikeIt.docx')
print("Document saved as Style_Guide_HikeIt.docx")
